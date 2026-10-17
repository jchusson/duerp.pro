import io
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas as reportlab_canvas
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

LOGO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "logo-jhc.png")

RED = colors.HexColor("#8B1D1D")
BLUE = colors.HexColor("#1B2F5E")
CREAM = colors.HexColor("#F5F0E8")
GREY = colors.HexColor("#D4D4C4")
GREEN = colors.HexColor("#2ECC71")
ORANGE = colors.HexColor("#F39C12")
RED_CELL = colors.HexColor("#E74C3C")
BLACK = colors.HexColor("#1F2937")

FREQUENCIES = [
    (1, "Faible : entre 0 et 10% du nombre d'heures travaillées"),
    (2, "Ponctuelle : entre 10 et 40% du nombre d'heures travaillées"),
    (3, "Moyenne : entre 40 et 70% du nombre d'heures travaillées"),
    (4, "Élevée : entre 70 et 100% du nombre d'heures travaillées"),
]
SEVERITIES = [
    (1, "Mineure : AT sans arrêt de travail"),
    (2, "Significative : AT avec arrêt de travail sans séquelles"),
    (4, "Critique : AT entraînant un arrêt de travail avec séquelles"),
    (8, "Vitale : AT pouvant entraîner une invalidité totale ou le décès"),
]
MASTERIES = [
    (0.5, "Bonne : personne protégée (protections individuelles, collectives, ...)"),
    (1, "Moyenne : risque signalé (informations, consignes, pas d'AT antérieur, ...)"),
    (2, "Aucune : aucune maîtrise"),
]


def _esc(t):
    return (t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def risk_color(value):
    if value <= 3:
        return GREEN
    if value < 16:
        return ORANGE
    return RED_CELL


def _header_footer(canvas, doc, client_name, doc_title):
    canvas.saveState()
    w, h = landscape(A4)
    canvas.setFillColor(BLUE)
    canvas.rect(0, h - 12 * mm, w, 12 * mm, stroke=0, fill=1)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 9)
    canvas.drawString(10 * mm, h - 8 * mm, "DUERP.pro — JHC Consulting")
    canvas.drawRightString(w - 10 * mm, h - 8 * mm, client_name)
    canvas.setFillColor(BLACK)
    canvas.setFont("Helvetica", 8)
    canvas.drawString(10 * mm, 6 * mm, doc_title)
    canvas.restoreState()


def _styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=base["Title"], fontSize=22, textColor=BLUE, alignment=1),
        "subtitle": ParagraphStyle("st", parent=base["Title"], fontSize=13, textColor=BLACK, alignment=1),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontSize=14, textColor=BLUE, spaceBefore=10),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontSize=12, textColor=RED, spaceBefore=8),
        "normal": ParagraphStyle("n", parent=base["Normal"], fontSize=9, leading=13),
        "cell": ParagraphStyle("c", parent=base["Normal"], fontSize=8, leading=10),
        "cellw": ParagraphStyle("cw", parent=base["Normal"], fontSize=8, leading=10, textColor=colors.white),
    }


LEGAL_TEXT = [
    ("Art. L. 4121-1", "L'employeur prend les mesures nécessaires pour assurer la sécurité et protéger la santé "
     "physique et mentale des travailleurs. Ces mesures comprennent : 1) Des actions de prévention des risques "
     "professionnels ; 2) Des actions d'information et de formation ; 3) La mise en place d'une organisation et de "
     "moyens adaptés. L'employeur veille à l'adaptation de ces mesures pour tenir compte du changement des "
     "circonstances et tendre à l'amélioration des situations existantes."),
    ("Art. L. 4121-3", "L'employeur, compte tenu de la nature des activités de l'établissement, évalue les risques "
     "pour la santé et la sécurité des travailleurs, y compris dans le choix des procédés de fabrication, des "
     "équipements de travail, des substances ou préparations chimiques, dans l'aménagement ou le réaménagement des "
     "lieux de travail ou des installations, dans l'organisation du travail et dans la définition des postes de "
     "travail."),
    ("Art. L. 4121-3-1", "Le document unique d'évaluation des risques professionnels répertorie l'ensemble des "
     "risques professionnels auxquels sont exposés les travailleurs et assure la traçabilité collective de ces "
     "expositions. L'employeur transcrit et met à jour dans le document unique les résultats de l'évaluation des "
     "risques pour la santé et la sécurité des travailleurs à laquelle il procède en application de l'article "
     "L. 4121-3."),
]

PREVENTION_PRINCIPLES = [
    "Éviter les risques",
    "Évaluer les risques qui ne peuvent pas être évités",
    "Combattre les risques à la source",
    "Adapter le travail à l'homme, en particulier en ce qui concerne la conception des postes de travail ainsi "
    "que le choix des équipements de travail et des méthodes de travail et de production, en vue notamment de "
    "limiter le travail monotone et le travail cadencé et de réduire les effets de ceux-ci sur la santé",
    "Tenir compte de l'état d'évolution de la technique",
    "Remplacer ce qui est dangereux par ce qui n'est pas dangereux ou par ce qui est moins dangereux",
    "Planifier la prévention en y intégrant, dans un ensemble cohérent, la technique, l'organisation du travail, "
    "les conditions de travail, les relations sociales et l'influence des facteurs ambiants, notamment les "
    "risques liés au harcèlement moral et au harcèlement sexuel",
    "Prendre des mesures de protection collective en leur donnant la priorité sur les mesures de protection "
    "individuelle",
    "Donner les instructions appropriées aux travailleurs",
]

METHODO_STEPS = [
    "Mise en place d'un groupe de travail avec représentativité de l'employeur.",
    "Sensibilisation de l'encadrement sur l'évaluation des risques professionnels et la démarche de prévention associée.",
    "Détermination des unités de travail (UT).",
    "Identification des risques par unité de travail a priori (sur le papier).",
    "Analyse des risques (Fréquence, Gravité, Maîtrise) des UT sur le terrain.",
    "Cotation et hiérarchisation des risques et conception du Document Unique.",
    "Mise en place d'un plan d'action.",
    "Validation du groupe de travail sur la totalité de la démarche et le Document Unique.",
    "Réévaluation périodique du Document Unique à la charge de l'employeur.",
]

RISK_LEVELS = [
    ("≤ 3", "AUCUNE ACTION complémentaire ne s'impose car l'événement dangereux a été pris en compte, étudié lors de "
     "l'évaluation des risques et la conclusion amène à un risque minime ou correctement maîtrisé par les mesures "
     "actuelles : tout nouveau personnel au poste concerné doit cependant être informé de l'éventualité de cet "
     "événement lors de la formation sécurité."),
    ("4 à 15", "Ce niveau de risque implique UNE ACTION, au moins en terme de communication et formalisation d'une "
     "consigne, de sensibilisation du personnel et de suivi dans le temps de l'évolution du risque."),
    ("≥ 16", "Ce niveau de risque implique UNE ACTION PALLIATIVE rapide et une SOLUTION DURABLE pour ramener le "
     "risque à un niveau acceptable."),
]


def build_pdf(client, units, actions, buffer):
    styles = _styles()
    today = date.today().strftime("%d/%m/%Y")
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=20 * mm, bottomMargin=12 * mm,
        title="Document Unique — %s" % client["name"],
    )
    story = []

    # Page de garde
    logos = []
    client_logo = client.get("logo_abs")
    if client_logo and os.path.exists(client_logo):
        logos.append(Image(client_logo, width=32 * mm, height=32 * mm))
    if os.path.exists(LOGO_PATH):
        logos.append(Image(LOGO_PATH, width=20 * mm, height=20 * mm))
    if logos:
        logo_tbl = Table([logos], colWidths=[40 * mm] * len(logos))
        logo_tbl.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "LEFT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story.append(logo_tbl)
    story.append(Spacer(1, 10 * mm))
    story.append(Paragraph(client["name"].upper(), styles["title"]))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("DOCUMENT UNIQUE", styles["h1"]))
    story.append(Paragraph("DE TRANSCRIPTION DES RÉSULTATS DE L'ÉVALUATION DES RISQUES PROFESSIONNELS", styles["subtitle"]))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("Établi en réponse aux exigences des articles L. 4121-1 à L. 4121-3 du Code du Travail.", styles["normal"]))
    story.append(Paragraph("Mise à jour le %s" % today, styles["normal"]))
    story.append(Spacer(1, 10 * mm))

    # Trois cadres de visa
    visa_titles = ["Visa du Chef d'établissement", "Visa du Rédacteur", "Visa du CSE / CST"]
    visa_cells = []
    for t in visa_titles:
        visa_cells.append([Paragraph("<b>%s</b>" % t, styles["cell"]), Spacer(1, 16 * mm), Paragraph("Nom :", styles["cell"]), Paragraph("Date :", styles["cell"]), Paragraph("Signature :", styles["cell"])])
    visa_tbl = Table([visa_cells], colWidths=[80 * mm, 80 * mm, 80 * mm])
    visa_tbl.setStyle(TableStyle([
        ("BOX", (0, 0), (0, 0), 1, BLACK),
        ("BOX", (1, 0), (1, 0), 1, BLACK),
        ("BOX", (2, 0), (2, 0), 1, BLACK),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(visa_tbl)
    story.append(PageBreak())

    # Fiche d'identité de l'établissement
    info = client.get("info") or {}
    story.append(Paragraph("Fiche d'identité de l'établissement", styles["h1"]))
    story.append(Spacer(1, 3 * mm))
    def _row(label, val):
        return Paragraph("<b>%s</b>" % label, styles["cell"]), Paragraph(str(val or ""), styles["cell"])
    rows = [
        _row("Raison sociale", info.get("raison_sociale") or client["name"]),
        _row("SIRET", info.get("siret")),
        _row("Téléphone", info.get("telephone")),
        _row("Adresse", info.get("adresse")),
        _row("Convention collective", info.get("convention_collective")),
        _row("Instance de représentation", info.get("instance_representa")),
    ]
    eff = info.get("cdi", 0) or 0
    total = sum(int(info.get(k) or 0) for k in ("cdi", "cdd", "interim", "alternance", "externe"))
    rows += [
        _row("Effectifs CDI / CDD / Intérim / Alternance / Externe", "%s / %s / %s / %s / %s" % (
            info.get("cdi", 0), info.get("cdd", 0), info.get("interim", 0), info.get("alternance", 0), info.get("externe", 0))),
        _row("Total effectif", total),
        _row("Dont Hommes / Femmes / RQTH", "%s / %s / %s" % (info.get("hommes", 0), info.get("femmes", 0), info.get("rqth", 0))),
        _row("Médecine du travail — organisme", info.get("med_organisme")),
        _row("Médecin du travail", info.get("med_medecin")),
        _row("Téléphone médecine", info.get("med_tel")),
        _row("Adresse médecine", info.get("med_adresse")),
        _row("Inspection du travail — inspecteur", info.get("it_inspecteur")),
        _row("Téléphone inspection", info.get("it_tel")),
        _row("Adresse inspection", info.get("it_adresse")),
    ]
    def _yn(v):
        return "Oui" if v else "Non"
    rows += [
        _row("Salle de repos / Réfectoire / Restaurant / Distributeur", "%s / %s / %s / %s" % (
            _yn(info.get("salle_repos")), _yn(info.get("refectoire")), _yn(info.get("restaurant")), _yn(info.get("distributeur")))),
        _row("Caractéristiques des locaux", info.get("locaux")),
        _row("Aération / Ventilation / Climatisation / Chauffage", "%s / %s / %s / %s" % (
            _yn(info.get("aeration")), _yn(info.get("ventilation")), _yn(info.get("clim")), _yn(info.get("chauffage")))),
        _row("Vestiaires / Douches / Toilettes", "%s / %s / %s" % (_yn(info.get("vestiaires")), _yn(info.get("douches")), _yn(info.get("toilettes")))),
        _row("Autres (hygiène)", info.get("hygiene_autres")),
        _row("Jours de travail", info.get("jours_travail")),
        _row("Temps de travail", info.get("temps_travail")),
        _row("Travail de nuit / posté / alternance", "%s / %s / %s" % (_yn(info.get("travail_nuit")), _yn(info.get("travail_poste")), _yn(info.get("travail_alternance")))),
        _row("AT avec arrêt / avec incapacité permanente", "%s / %s" % (info.get("at_avec_arret", 0), info.get("at_incapacite", 0))),
        _row("Nombre d'enquêtes sur les accidents", info.get("nb_enquetes", 0)),
        _row("Commentaires", info.get("at_commentaires")),
        _row("Maladies à caractère professionnel (médecin / autres sources)", "%s / %s" % (info.get("mcp_medecin", 0), info.get("mcp_autres", 0))),
        _row("Nature", info.get("mcp_nature")),
        _row("Maladies professionnelles (médecin du travail / autre médecin)", "%s / %s" % (info.get("mp_medecin", 0), info.get("mp_autre_medecin", 0))),
    ]
    fiche_tbl = Table(rows, colWidths=[95 * mm, 150 * mm])
    fiche_tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (0, -1), CREAM),
    ]))
    story.append(fiche_tbl)
    story.append(PageBreak())

    # Cadre légal
    story.append(Paragraph("Cadre légal", styles["h1"]))
    for title, text in LEGAL_TEXT:
        story.append(Paragraph(title, styles["h2"]))
        story.append(Paragraph(text, styles["normal"]))
    story.append(Paragraph("Art. L. 4121-2 — Les 9 principes généraux de prévention", styles["h2"]))
    story.append(Paragraph(
        "L'employeur met en œuvre les mesures prévues à l'article L. 4121-1 sur le fondement des principes "
        "généraux de prévention suivants :", styles["normal"]))
    principles = [[Paragraph("<b>%d.</b> %s" % (i, p), styles["normal"])] for i, p in enumerate(PREVENTION_PRINCIPLES, 1)]
    story.append(Table(principles, colWidths=[245 * mm], hAlign="LEFT"))
    story.append(PageBreak())

    # Méthodologie
    story.append(Paragraph("I) Principe de la méthodologie d'évaluation des risques utilisée", styles["h1"]))
    story.append(Paragraph(
        "Cette démarche a pour objectif de permettre de prendre les mesures nécessaires pour assurer la sécurité "
        "et la santé de ses salariés au travail.", styles["normal"]))
    for i, step in enumerate(METHODO_STEPS, 1):
        story.append(Paragraph("%d. %s" % (i, step), styles["normal"]))
    story.append(Paragraph("II) Grille de cotation des risques", styles["h2"]))
    story.append(Paragraph("Fréquences d'expositions", styles["cell"]))
    for v, label in FREQUENCIES:
        story.append(Paragraph("%d : %s" % (v, label), styles["cell"]))
    story.append(Paragraph("Gravités des risques", styles["cell"]))
    for v, label in SEVERITIES:
        story.append(Paragraph("%d : %s" % (v, label), styles["cell"]))
    story.append(Paragraph("Maîtrises des risques", styles["cell"]))
    for v, label in MASTERIES:
        story.append(Paragraph("%s : %s" % (("%g" % v).replace(".", ","), label), styles["cell"]))
    story.append(Paragraph(
        "La cotation d'évaluation du risque est obtenue en multipliant la fréquence × la gravité × la maîtrise.",
        styles["normal"]))
    story.append(Paragraph("III) Niveau de risque", styles["h2"]))
    lvl_rows = [[Paragraph("Niveau de risque", styles["cellw"]), Paragraph("Signification", styles["cellw"])]]
    for lvl, txt in RISK_LEVELS:
        lvl_rows.append([Paragraph(lvl, styles["cell"]), Paragraph(txt, styles["cell"])])
    lvl_tbl = Table(lvl_rows, colWidths=[35 * mm, 230 * mm])
    lvl_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), RED),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(lvl_tbl)
    story.append(Paragraph("IV) Listing des unités de travail", styles["h2"]))
    for u in units:
        story.append(Paragraph(u["name"], styles["normal"]))
    story.append(PageBreak())

    # Tableaux par unité de travail
    col_w = [38 * mm, 82 * mm, 12 * mm, 12 * mm, 82 * mm, 14 * mm, 16 * mm, 44 * mm]
    for u in units:
        risks = u["risks"]
        if not risks:
            continue
        story.append(Paragraph("Unité de travail : %s" % u["name"], styles["h1"]))
        story.append(Paragraph("Généré le %s" % today, styles["cell"]))
        header = [Paragraph(h, styles["cellw"]) for h in
                  ["Risque", "Situation dangereuse", "F", "G", "Mesures de prévention", "M", "FxGxM", "Actions de prévention"]]
        rows = [header]
        cell_style = []
        for idx, r in enumerate(risks, start=1):
            crit = round(r["frequency"] * r["severity"] * r["mastery"], 1)
            cell_style.append(("BACKGROUND", (6, idx), (6, idx), risk_color(crit)))
            measures_txt = "<br/>".join("&#8226; " + _esc(m["text"]) for m in r["measures"])
            def _action_line(a):
                meta = []
                if a.get("actor"):
                    meta.append(_esc(a["actor"]))
                if a.get("due_date"):
                    meta.append(_esc(a["due_date"]))
                if meta:
                    return "&#8226; %s <font size=6 color=#78716c>(%s)</font>" % (_esc(a["text"]), " — ".join(meta))
                return "&#8226; " + _esc(a["text"])
            actions_txt = "<br/>".join(_action_line(a) for a in r["actions"])
            rows.append([
                Paragraph(r["hazard"] or "", styles["cell"]),
                Paragraph(r["situation"] or "", styles["cell"]),
                Paragraph(str(r["frequency"]), styles["cell"]),
                Paragraph(str(r["severity"]), styles["cell"]),
                Paragraph(measures_txt, styles["cell"]),
                Paragraph(("%g" % r["mastery"]).replace(".", ","), styles["cell"]),
                Paragraph(("%g" % crit).replace(".", ","), styles["cell"]),
                Paragraph(actions_txt, styles["cell"]),
            ])
        tbl = Table(rows, colWidths=col_w, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), GREY),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (2, 0), (6, -1), "CENTER"),
        ] + cell_style))
        story.append(tbl)
        story.append(PageBreak())

    # Plan d'action
    story.extend(_action_flowables(actions, styles, today))

    doc.build(
        story,
        onFirstPage=lambda c, d: _header_footer(c, d, client["name"], "Document Unique"),
        onLaterPages=lambda c, d: _header_footer(c, d, client["name"], "Document Unique"),
        canvasmaker=_NumberedCanvas,
    )


class _NumberedCanvas(reportlab_canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.setFont("Helvetica", 8)
            self.drawRightString(self._pagesize[0] - 15 * mm, 6 * mm,
                                  "%d / %d" % (self._pageNumber, total))
            super().showPage()
        super().save()


def _action_flowables(actions, styles, today):
    flow = []
    flow.append(Paragraph("Plan d'action", styles["h1"]))
    flow.append(Paragraph("Généré le %s" % today, styles["cell"]))
    if not actions:
        flow.append(Paragraph("Aucune action enregistrée.", styles["normal"]))
        return flow
    header = [Paragraph(h, styles["cellw"]) for h in
              ["Unité de travail", "Action", "Priorité", "Responsable", "Échéance", "Statut", "Coût estimé"]]
    status_colors = {"À faire": ORANGE, "En cours": BLUE, "Terminée": GREEN}
    rows = [header]
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), GREY),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for idx, a in enumerate(actions, start=1):
        rows.append([
            Paragraph(a["work_unit"] or "", styles["cell"]),
            Paragraph(a["action"], styles["cell"]),
            Paragraph(a["priority"], styles["cell"]),
            Paragraph(a["responsible"], styles["cell"]),
            Paragraph(a["due_date"] or "", styles["cell"]),
            Paragraph(a["status"], styles["cell"]),
            Paragraph(a["cost"], styles["cell"]),
        ])
        if a["status"] in status_colors:
            style_cmds.append(("TEXTCOLOR", (5, idx), (5, idx), status_colors[a["status"]]))
    tbl = Table(rows, colWidths=[50 * mm, 90 * mm, 22 * mm, 40 * mm, 26 * mm, 22 * mm, 28 * mm], repeatRows=1)
    tbl.setStyle(TableStyle(style_cmds))
    flow.append(tbl)
    return flow


def build_action_pdf(client, actions, buffer):
    styles = _styles()
    today = date.today().strftime("%d/%m/%Y")
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=15 * mm, rightMargin=15 * mm, topMargin=18 * mm, bottomMargin=14 * mm,
        title="Plan d'action — %s" % client["name"],
    )
    story = []
    if os.path.exists(LOGO_PATH):
        story.append(Image(LOGO_PATH, width=22 * mm, height=22 * mm))
    story.extend(_action_flowables(actions, styles, today))
    doc.build(
        story,
        onFirstPage=lambda c, d: _header_footer(c, d, client["name"], "Plan d'action"),
        onLaterPages=lambda c, d: _header_footer(c, d, client["name"], "Plan d'action"),
        canvasmaker=_NumberedCanvas,
    )


def build_excel(client, units, actions, buffer):
    wb = Workbook()
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D4D4C4")
    red_fill = PatternFill("solid", fgColor="8B1D1D")
    header_font = Font(bold=True)

    ws = wb.active
    ws.title = "Document Unique"
    ws["A1"] = client["name"]
    ws["A1"].font = Font(bold=True, size=16, color="1B2F5E")
    ws["A2"] = "DOCUMENT UNIQUE DE TRANSCRIPTION DES RÉSULTATS DE L'ÉVALUATION DES RISQUES PROFESSIONNELS"
    ws["A3"] = "Mise à jour le %s" % date.today().strftime("%d/%m/%Y")
    headers = ["Unité de travail", "Risque", "Situation dangereuse", "F", "G",
               "Mesures de prévention", "M", "FxGxM", "Actions de prévention"]
    ws.append([])
    ws.append(headers)
    hrow = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=hrow, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    for u in units:
        for r in u["risks"]:
            crit = round(r["frequency"] * r["severity"] * r["mastery"], 1)
            ws.append([u["name"], r["hazard"], r["situation"], r["frequency"], r["severity"],
                       "\n".join(m["text"] for m in r["measures"]), r["mastery"], crit,
                       "\n".join(a["text"] + ((" (%s)" % " — ".join(filter(None, [a["actor"], a["due_date"]]))) if (a["actor"] or a["due_date"]) else "") for a in r["actions"])])
            row = ws.max_row
            color = "2ECC71" if crit <= 3 else ("F39C12" if crit < 16 else "E74C3C")
            ws.cell(row=row, column=8).fill = PatternFill("solid", fgColor=color)
            for c in range(1, len(headers) + 1):
                ws.cell(row=row, column=c).border = border
    widths = [30, 22, 40, 5, 5, 40, 6, 8, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A%d" % (hrow + 1)

    wa = wb.create_sheet("Plan d'action")
    ah = ["Unité de travail", "Action", "Priorité", "Responsable", "Échéance", "Statut", "Coût estimé"]
    wa.append(ah)
    for c in range(1, len(ah) + 1):
        cell = wa.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    for a in actions:
        wa.append([a["work_unit"], a["action"], a["priority"], a["responsible"], a["due_date"], a["status"], a["cost"]])
        for c in range(1, len(ah) + 1):
            wa.cell(row=wa.max_row, column=c).border = border
    for i, w in enumerate([30, 50, 12, 20, 14, 12, 14], 1):
        wa.column_dimensions[get_column_letter(i)].width = w

    wm = wb.create_sheet("Méthodologie")
    wm.append(["Fréquences d'exposition"])
    for v, label in FREQUENCIES:
        wm.append([v, label])
    wm.append([])
    wm.append(["Gravités des risques"])
    for v, label in SEVERITIES:
        wm.append([v, label])
    wm.append([])
    wm.append(["Maîtrises des risques"])
    for v, label in MASTERIES:
        wm.append([v, label])
    wm.append([])
    wm.append(["Niveau de risque", "Signification"])
    for lvl, txt in RISK_LEVELS:
        wm.append([lvl, txt])
    wm.column_dimensions["A"].width = 22
    wm.column_dimensions["B"].width = 110
    wm["A1"].font = header_font
    wb.save(buffer)


def build_action_excel(client, actions, buffer):
    wb = Workbook()
    wa = wb.active
    wa.title = "Plan d'action"
    wa["A1"] = client["name"]
    wa["A1"].font = Font(bold=True, size=14, color="1B2F5E")
    wa["A2"] = "Plan d'action — mise à jour le %s" % date.today().strftime("%d/%m/%Y")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D4D4C4")
    ah = ["Unité de travail", "Action", "Priorité", "Responsable", "Échéance", "Statut", "Coût estimé"]
    wa.append([])
    wa.append(ah)
    hrow = wa.max_row
    for c in range(1, len(ah) + 1):
        cell = wa.cell(row=hrow, column=c)
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.border = border
    for a in actions:
        wa.append([a["work_unit"], a["action"], a["priority"], a["responsible"], a["due_date"], a["status"], a["cost"]])
        for c in range(1, len(ah) + 1):
            wa.cell(row=wa.max_row, column=c).border = border
    for i, w in enumerate([30, 50, 12, 20, 14, 12, 14], 1):
        wa.column_dimensions[get_column_letter(i)].width = w
    wa.freeze_panes = "A%d" % (hrow + 1)
    wb.save(buffer)
