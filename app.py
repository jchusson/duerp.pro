import io
import os
import sys
import tempfile
import time

from flask import (
    Flask, abort, flash, g, jsonify, redirect, render_template,
    request, session, send_file, url_for,
)
from werkzeug.security import generate_password_hash

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import database
import exports
import ai

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", os.urandom(24).hex())

FREQ = [1, 2, 3, 4]
SEV = [1, 2, 4, 8]
MASTERY = [0.5, 1, 2]


@app.before_request
def before_request():
    database.init_db()
    g.user = None
    user_id = session.get("user_id")
    if user_id:
        g.user = database.get_user(user_id)


def login_required(f):
    def wrapper(*args, **kwargs):
        if not g.get("user"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    wrapper.__name__ = f.__name__
    return wrapper


def admin_required(f):
    def wrapper(*args, **kwargs):
        if not g.get("user") or g.user["role"] != "admin":
            abort(403)
        return f(*args, **kwargs)
    wrapper.__name__ = f.__name__
    return wrapper


def current_client():
    user = g.get("user")
    if not user:
        return None
    client_id = session.get("client_id")
    clients = database.user_clients(user)
    if not clients:
        return None
    if client_id:
        for c in clients:
            if c["id"] == client_id:
                return c
    session["client_id"] = clients[0]["id"]
    return clients[0]


# ---------------- Auth ----------------

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = database.verify_user(request.form.get("email", "").strip().lower(), request.form.get("password", ""))
        if not user:
            flash("Identifiants incorrects.", "error")
        else:
            session.clear()
            session["user_id"] = user["id"]
            return redirect(url_for("home"))
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


def _load_units(client_id):
    """Charge les UT avec risques, mesures et actions de prévention."""
    conn = database.get_db()
    units = [dict(u) for u in conn.execute(
        "SELECT * FROM work_units WHERE client_id = ? ORDER BY position, id", (client_id,)
    ).fetchall()]
    for u in units:
        u["risks"] = [dict(r) for r in conn.execute(
            "SELECT * FROM risks WHERE work_unit_id = ? ORDER BY id", (u["id"],)
        ).fetchall()]
        for r in u["risks"]:
            r["measures"] = [dict(m) for m in conn.execute(
                "SELECT * FROM measures WHERE risk_id = ? ORDER BY id", (r["id"],)
            ).fetchall()]
            r["actions"] = [dict(m) for m in conn.execute(
                "SELECT * FROM actions WHERE risk_id = ? ORDER BY id", (r["id"],)
            ).fetchall()]
    conn.close()
    return units


def crit_value(r):
    return r["frequency"] * r["severity"] * r["mastery"]


@app.route("/")
@login_required
def home():
    client = current_client()
    if not client:
        if g.user["role"] == "admin":
            return redirect(url_for("admin"))
        return render_template("no_client.html")
    q = request.args.get("q", "").strip().lower()
    crit_level = request.args.get("crit", "")
    unit_filter = request.args.get("unit", "")
    units = _load_units(client["id"])
    count_risks = sum(len(u["risks"]) for u in units)
    conn = database.get_db()
    actions_plan = conn.execute(
        "SELECT a.*, w.name AS work_unit FROM action_items a "
        "LEFT JOIN work_units w ON w.id = a.work_unit_id "
        "WHERE a.client_id = ? ORDER BY a.id DESC", (client["id"],)
    ).fetchall()
    conn.close()

    # --- filtres ---
    units_shown = units
    filters_active = bool(q or crit_level or unit_filter)
    if filters_active:
        try:
            unit_id_int = int(unit_filter) if unit_filter else None
        except ValueError:
            unit_id_int = None
        for u in units_shown:
            if unit_id_int and u["id"] != unit_id_int:
                u["risks"] = []
                continue
            kept = []
            for r in u["risks"]:
                if q:
                    texts = [u["name"], r["hazard"], r["situation"],
                             " ".join(m["text"] for m in r["measures"]),
                             " ".join(a["text"] for a in r["actions"])]
                    if q not in " ".join(texts).lower():
                        continue
                v = crit_value(r)
                if crit_level == "low" and v > 3:
                    continue
                if crit_level == "mid" and not (4 <= v <= 15):
                    continue
                if crit_level == "high" and v < 16:
                    continue
                kept.append(r)
            u["risks"] = kept
        units_shown = [u for u in units_shown if u["risks"]]

    high = sum(1 for u in units for r in u["risks"] if crit_value(r) >= 16)

    conn = database.get_db()
    last_update = _last_update(conn, client["id"])
    last_archive = conn.execute(
        "SELECT created_at, filename FROM archives WHERE client_id=? "
        "ORDER BY id DESC LIMIT 1", (client["id"],)
    ).fetchone()
    archives = conn.execute(
        "SELECT * FROM archives WHERE client_id=? ORDER BY id DESC LIMIT 5", (client["id"],)
    ).fetchall()
    ai_client_enabled = bool(conn.execute(
        "SELECT ai_enabled FROM clients WHERE id=?", (client["id"],)).fetchone()[0])
    conn.close()
    from risk_catalog import RISQUE_CATALOG
    return render_template(
        "home.html", client=client, units=units, actions=actions_plan,
        count_risks=count_risks, count_high=high,
        units_shown=units_shown,
        filters_active=filters_active,
        catalog=RISQUE_CATALOG,
        ai_enabled=ai.is_configured() and ai_client_enabled,
        ai_client_enabled=ai_client_enabled,
        last_update=last_update,
        last_archive=dict(last_archive) if last_archive else None,
        archives=[dict(x) for x in archives],
        clients=database.user_clients(g.user),
    )


@app.route("/client/select/<int:client_id>", methods=["POST"])
@login_required
def select_client(client_id):
    if not database.user_has_client(g.user, client_id):
        abort(403)
    session["client_id"] = client_id
    return redirect(url_for("home"))


# ---------------- Unités de travail ----------------

@app.route("/work-units/add", methods=["POST"])
@login_required
def add_work_unit():
    client = current_client()
    if not client:
        abort(403)
    name = request.form.get("name", "").strip()
    if name:
        conn = database.get_db()
        conn.execute(
            "INSERT INTO work_units (client_id, name, position) VALUES (?,?,?)",
            (client["id"], name, request.form.get("position", 0) or 0),
        )
        conn.commit()
        conn.close()
    return redirect(url_for("home"))


@app.route("/work-units/<int:uid>/rename", methods=["POST"])
@login_required
def rename_work_unit(uid):
    client = current_client()
    if not client:
        abort(403)
    name = request.form.get("name", "").strip()
    if name:
        conn = database.get_db()
        conn.execute(
            "UPDATE work_units SET name = ? WHERE id = ? AND client_id = ?",
            (name, uid, client["id"]),
        )
        conn.commit()
        conn.close()
    return redirect(url_for("home"))


@app.route("/work-units/duplicate", methods=["POST"])
@login_required
def duplicate_work_unit_selected():
    """Formulaire « Dupliquer une UT » (sélecteur src)."""
    try:
        uid = int(request.form.get("src", ""))
    except (TypeError, ValueError):
        return redirect(url_for("home"))
    return duplicate_work_unit(uid)


@app.route("/work-units/<int:uid>/duplicate", methods=["POST"])
@login_required
def duplicate_work_unit(uid):
    client = current_client()
    if not client:
        abort(403)
    conn = database.get_db()
    src = conn.execute(
        "SELECT * FROM work_units WHERE id = ? AND client_id = ?", (uid, client["id"])
    ).fetchone()
    if not src:
        conn.close()
        abort(403)
    import re as _re
    base_name = _re.sub(r"\s*\(copie(\s*\d+)?\)$", "", src["name"]).strip()
    if not conn.execute(
        "SELECT id FROM work_units WHERE client_id=? AND name=?",
        (client["id"], base_name + " (copie)"),
    ).fetchone():
        new_name = base_name + " (copie)"
    else:
        i = 2
        while conn.execute(
            "SELECT id FROM work_units WHERE client_id=? AND name=?",
            (client["id"], "%s (copie %d)" % (base_name, i))).fetchone():
            i += 1
        new_name = "%s (copie %d)" % (base_name, i)
    cur = conn.execute(
        "INSERT INTO work_units (client_id, name, position) VALUES (?,?,?)",
        (client["id"], new_name, src["position"]),
    )
    new_uid = cur.lastrowid
    risks = conn.execute("SELECT * FROM risks WHERE work_unit_id = ?", (uid,)).fetchall()
    for r in risks:
        cur = conn.execute(
            "INSERT INTO risks (work_unit_id, hazard, situation, frequency, severity, mastery) "
            "VALUES (?,?,?,?,?,?)",
            (new_uid, r["hazard"], r["situation"], r["frequency"], r["severity"], r["mastery"]),
        )
        new_rid = cur.lastrowid
        for m in conn.execute("SELECT * FROM measures WHERE risk_id = ?", (r["id"],)).fetchall():
            conn.execute("INSERT INTO measures (risk_id, text) VALUES (?,?)", (new_rid, m["text"]))
        for a in conn.execute("SELECT * FROM actions WHERE risk_id = ?", (r["id"],)).fetchall():
            conn.execute("INSERT INTO actions (risk_id, text) VALUES (?,?)", (new_rid, a["text"]))
    conn.commit()
    conn.close()
    return redirect(url_for("home"))


@app.route("/work-units/<int:uid>/delete", methods=["POST"])
@login_required
def delete_work_unit(uid):
    client = current_client()
    if not client:
        abort(403)
    conn = database.get_db()
    conn.execute(
        "DELETE FROM work_units WHERE id = ? AND client_id = ?", (uid, client["id"])
    )
    conn.commit()
    conn.close()
    return redirect(url_for("home"))


# ---------------- Risques ----------------

def _risk_state(conn, rid):
    """État complet d'une situation dangereuse (pour les réponses JSON)."""
    r = conn.execute("SELECT * FROM risks WHERE id=?", (rid,)).fetchone()
    if not r:
        return None
    return {
        "id": r["id"], "hazard": r["hazard"], "situation": r["situation"],
        "frequency": r["frequency"], "severity": r["severity"], "mastery": r["mastery"],
        "measures": [dict(m) for m in conn.execute(
            "SELECT id, text FROM measures WHERE risk_id=? ORDER BY id", (rid,)).fetchall()],
        "actions": [dict(x) for x in conn.execute(
            "SELECT id, text, actor, due_date FROM actions WHERE risk_id=? ORDER BY id", (rid,)).fetchall()],
    }


def _is_ajax():
    return request.headers.get("X-Requested-With") == "fetch"


def _valid_risk_params():
    try:
        frequency = int(request.form.get("frequency"))
        severity = int(request.form.get("severity"))
        mastery = float(request.form.get("mastery"))
    except (TypeError, ValueError):
        abort(400)
    if frequency not in FREQ or severity not in SEV or mastery not in MASTERY:
        abort(400)
    return frequency, severity, mastery


@app.route("/risks/add", methods=["POST"])
@login_required
def add_risk():
    client = current_client()
    if not client:
        abort(403)
    try:
        work_unit_id = int(request.form.get("work_unit_id"))
    except (TypeError, ValueError):
        abort(400)
    frequency, severity, mastery = _valid_risk_params()
    conn = database.get_db()
    check = conn.execute(
        "SELECT id FROM work_units WHERE id = ? AND client_id = ?", (work_unit_id, client["id"])
    ).fetchone()
    if not check:
        conn.close()
        abort(403)
    cur = conn.execute(
        "INSERT INTO risks (work_unit_id, hazard, situation, frequency, severity, mastery) "
        "VALUES (?,?,?,?,?,?)",
        (work_unit_id, request.form.get("hazard", "").strip(),
         request.form.get("situation", "").strip(), frequency, severity, mastery),
    )
    new_rid = cur.lastrowid
    for m in request.form.getlist("measures[]"):
        if m.strip():
            conn.execute("INSERT INTO measures (risk_id, text) VALUES (?,?)", (new_rid, m.strip()))
    for a in request.form.getlist("actions[]"):
        if a.strip():
            conn.execute("INSERT INTO actions (risk_id, text) VALUES (?,?)", (new_rid, a.strip()))
    conn.commit()
    conn.close()
    return redirect(url_for("home"))


@app.route("/risks/<int:rid>/update", methods=["POST"])
@login_required
def update_risk(rid):
    client = current_client()
    if not client:
        abort(403)
    frequency, severity, mastery = _valid_risk_params()
    conn = database.get_db()
    row = conn.execute(
        "SELECT id FROM risks WHERE id=? AND work_unit_id IN "
        "(SELECT id FROM work_units WHERE client_id=?)", (rid, client["id"])
    ).fetchone()
    if not row:
        conn.close()
        abort(403)
    conn.execute(
        "UPDATE risks SET hazard=?, situation=?, frequency=?, severity=?, mastery=? WHERE id=?",
        (request.form.get("hazard", "").strip(), request.form.get("situation", "").strip(),
         frequency, severity, mastery, rid),
    )
    if "measures[]" in request.form:
        conn.execute("DELETE FROM measures WHERE risk_id=?", (rid,))
        for m in request.form.getlist("measures[]"):
            if m.strip():
                conn.execute("INSERT INTO measures (risk_id, text) VALUES (?,?)", (rid, m.strip()))
    if "actions[]" in request.form:
        conn.execute("DELETE FROM actions WHERE risk_id=?", (rid,))
        for a in request.form.getlist("actions[]"):
            if a.strip():
                conn.execute("INSERT INTO actions (risk_id, text) VALUES (?,?)", (rid, a.strip()))
    conn.commit()
    if _is_ajax():
        state = _risk_state(conn, rid)
        conn.close()
        return jsonify(state)
    conn.close()
    return redirect(url_for("home"))


@app.route("/risks/<int:rid>/state")
@login_required
def risk_state(rid):
    cl = current_client()
    if not cl:
        abort(403)
    conn = database.get_db()
    row = conn.execute(
        "SELECT id FROM risks WHERE id=? AND work_unit_id IN "
        "(SELECT id FROM work_units WHERE client_id=?)", (rid, cl["id"])
    ).fetchone()
    if not row:
        conn.close()
        abort(404)
    state = _risk_state(conn, rid)
    conn.close()
    return jsonify(state)


@app.route("/risks/<int:rid>/duplicate", methods=["POST"])
@login_required
def duplicate_risk(rid):
    """Duplique une situation dangereuse vers une ou plusieurs unités de travail."""
    client = current_client()
    if not client:
        abort(403)
    target_ids = [int(t) for t in request.form.getlist("targets[]") if t.strip().isdigit()]
    conn = database.get_db()
    src = conn.execute(
        "SELECT * FROM risks WHERE id=? AND work_unit_id IN "
        "(SELECT id FROM work_units WHERE client_id=?)", (rid, client["id"])
    ).fetchone()
    if src:
        measures = conn.execute("SELECT * FROM measures WHERE risk_id=?", (rid,)).fetchall()
        actions = conn.execute("SELECT * FROM actions WHERE risk_id=?", (rid,)).fetchall()
        for tid in target_ids:
            check = conn.execute(
                "SELECT id FROM work_units WHERE id=? AND client_id=?", (tid, client["id"])
            ).fetchone()
            if not check:
                continue
            cur = conn.execute(
                "INSERT INTO risks (work_unit_id, hazard, situation, frequency, severity, mastery) "
                "VALUES (?,?,?,?,?,?)",
                (tid, src["hazard"], src["situation"], src["frequency"], src["severity"], src["mastery"]),
            )
            new_rid = cur.lastrowid
            for m in measures:
                conn.execute("INSERT INTO measures (risk_id, text) VALUES (?,?)", (new_rid, m["text"]))
            for a in actions:
                conn.execute(
                    "INSERT INTO actions (risk_id, text, actor, due_date) VALUES (?,?,?,?)",
                    (new_rid, a["text"], a["actor"], a["due_date"]),
                )
        conn.commit()
    conn.close()
    return redirect(url_for("home"))


@app.route("/risks/<int:rid>/delete", methods=["POST"])
@login_required
def delete_risk(rid):
    client = current_client()
    if not client:
        abort(403)
    conn = database.get_db()
    conn.execute(
        "DELETE FROM risks WHERE id=? AND work_unit_id IN (SELECT id FROM work_units WHERE client_id=?)",
        (rid, client["id"]),
    )
    conn.commit()
    conn.close()
    return redirect(url_for("home"))


# ---------------- Mesures et actions de prévention (lignes individuelles) ----------------

@app.route("/risks/<int:rid>/measures/add", methods=["POST"])
@login_required
def add_measure(rid):
    return _add_line(rid, "measures")


@app.route("/risks/<int:rid>/actions/add", methods=["POST"])
@login_required
def add_action_line(rid):
    return _add_line(rid, "actions")


def _add_line(rid, table):
    client = current_client()
    if not client:
        abort(403)
    text = request.form.get("text", "").strip()
    if not text:
        if _is_ajax():
            return jsonify(error="texte vide"), 400
        return redirect(url_for("home"))
    conn = database.get_db()
    row = conn.execute(
        "SELECT id FROM risks WHERE id=? AND work_unit_id IN "
        "(SELECT id FROM work_units WHERE client_id=?)", (rid, client["id"])
    ).fetchone()
    if row:
        if table == "actions":
            conn.execute(
                "INSERT INTO actions (risk_id, text, actor, due_date) VALUES (?,?,?,?)",
                (rid, text,
                 request.form.get("actor", "").strip(),
                 request.form.get("due_date", "").strip()),
            )
        else:
            conn.execute("INSERT INTO measures (risk_id, text) VALUES (?,?)", (rid, text))
        conn.commit()
    if _is_ajax():
        state = _risk_state(conn, rid) if row else None
        conn.close()
        if not state:
            return jsonify(error="introuvable"), 404
        return jsonify(state)
    conn.close()
    return redirect(url_for("home"))


@app.route("/measures/<int:mid>/update", methods=["POST"])
@login_required
def update_measure(mid):
    return _update_line(mid, "measures")


@app.route("/actions/<int:aid>/update", methods=["POST"])
@login_required
def update_action_line(aid):
    return _update_line(aid, "actions")


def _update_line(line_id, table):
    client = current_client()
    if not client:
        abort(403)
    text = request.form.get("text", "").strip()
    conn = database.get_db()
    row = conn.execute(
        "SELECT l.id, l.risk_id FROM %s l JOIN risks r ON r.id = l.risk_id "
        "JOIN work_units w ON w.id = r.work_unit_id WHERE l.id = ? AND w.client_id = ?" % table,
        (line_id, client["id"]),
    ).fetchone()
    if row:
        if table == "actions":
            conn.execute(
                "UPDATE actions SET text = ?, actor = ?, due_date = ? WHERE id = ?",
                (text, request.form.get("actor", "").strip(),
                 request.form.get("due_date", "").strip(), line_id),
            )
        else:
            conn.execute("UPDATE measures SET text = ? WHERE id = ?", (text, line_id))
        conn.commit()
    if _is_ajax():
        state = _risk_state(conn, row["risk_id"]) if row else None
        conn.close()
        if not state:
            return jsonify(error="introuvable"), 404
        return jsonify(state)
    conn.close()
    return redirect(url_for("home"))


@app.route("/measures/<int:mid>/delete", methods=["POST"])
@login_required
def delete_measure(mid):
    return _delete_line(mid, "measures")


@app.route("/actions/<int:aid>/delete", methods=["POST"])
@login_required
def delete_action_line(aid):
    return _delete_line(aid, "actions")


def _delete_line(line_id, table):
    client = current_client()
    if not client:
        abort(403)
    conn = database.get_db()
    row = conn.execute(
        "SELECT l.risk_id FROM %s l JOIN risks r ON r.id = l.risk_id "
        "JOIN work_units w ON w.id = r.work_unit_id WHERE l.id = ? AND w.client_id = ?" % table,
        (line_id, client["id"]),
    ).fetchone()
    conn.execute(
        "DELETE FROM %s WHERE id = ?" % table, (line_id,)
    )
    conn.commit()
    if _is_ajax():
        state = _risk_state(conn, row["risk_id"]) if row else None
        conn.close()
        if not state:
            return jsonify(error="introuvable"), 404
        return jsonify(state)
    conn.close()
    return redirect(url_for("home"))


# ---------------- Plan d'action ----------------

@app.route("/actions/add", methods=["POST"])
@login_required
def add_action():
    client = current_client()
    if not client:
        abort(403)
    action = request.form.get("action", "").strip()
    if not action:
        return redirect(url_for("home"))
    work_unit_id = request.form.get("work_unit_id") or None
    conn = database.get_db()
    conn.execute(
        "INSERT INTO action_items (client_id, work_unit_id, action, priority, responsible, due_date, status, cost) "
        "VALUES (?,?,?,?,?,?,?,?)",
        (client["id"], work_unit_id, action,
         request.form.get("priority", "Moyenne"), request.form.get("responsible", "").strip(),
         request.form.get("due_date", "").strip(), request.form.get("status", "À faire"),
         request.form.get("cost", "").strip()),
    )
    conn.commit()
    conn.close()
    return redirect(url_for("home"))


@app.route("/actions/<int:aid>/update", methods=["POST"])
@login_required
def update_action(aid):
    client = current_client()
    if not client:
        abort(403)
    conn = database.get_db()
    conn.execute(
        "UPDATE action_items SET status=? WHERE id=? AND client_id=?",
        (request.form.get("status", "À faire"), aid, client["id"]),
    )
    conn.commit()
    conn.close()
    return redirect(url_for("home"))


@app.route("/actions/<int:aid>/delete", methods=["POST"])
@login_required
def delete_action(aid):
    client = current_client()
    if not client:
        abort(403)
    conn = database.get_db()
    conn.execute("DELETE FROM action_items WHERE id=? AND client_id=?", (aid, client["id"]))
    conn.commit()
    conn.close()
    return redirect(url_for("home"))


# ---------------- Fiche d'identité / logo / archives ----------------

INFO_FIELDS = [
    "raison_sociale", "siret", "telephone", "adresse", "convention_collective",
    "instance_representa",
]
INFO_INTS = [
    "cdi", "cdd", "interim", "alternance", "externe", "hommes", "femmes", "rqth",
    "salle_repos", "refectoire", "restaurant", "distributeur",
    "aeration", "ventilation", "clim", "chauffage", "vestiaires", "douches", "toilettes",
    "travail_nuit", "travail_poste", "travail_alternance",
    "at_avec_arret", "at_incapacite", "nb_enquetes",
    "mcp_medecin", "mcp_autres", "mp_medecin", "mp_autre_medecin",
]
INFO_TEXTS = [
    "med_organisme", "med_medecin", "med_tel", "med_adresse",
    "it_inspecteur", "it_tel", "it_adresse",
    "locaux", "hygiene_autres", "jours_travail", "temps_travail",
    "at_commentaires", "mcp_nature",
]


@app.route("/client/info", methods=["GET", "POST"])
@login_required
def client_info():
    cl = current_client()
    if not cl:
        abort(403)
    conn = database.get_db()
    if request.method == "POST":
        vals = {}
        for f in INFO_FIELDS + INFO_TEXTS:
            vals[f] = request.form.get(f, "").strip()
        for f in INFO_INTS:
            try:
                vals[f] = int(request.form.get(f, 0) or 0)
            except ValueError:
                vals[f] = 0
        sets = ", ".join("%s=?" % f for f in vals)
        exists = conn.execute("SELECT client_id FROM client_info WHERE client_id=?", (cl["id"],)).fetchone()
        if exists:
            conn.execute(
                "UPDATE client_info SET %s, updated_at=datetime('now','localtime') WHERE client_id=?" % sets,
                (*vals.values(), cl["id"]))
        else:
            conn.execute(
                "INSERT INTO client_info (client_id, %s) VALUES (?,%s)" % (
                    ", ".join(vals), ", ".join("?" * len(vals))),
                (cl["id"], *vals.values()))
        conn.commit()
        flash("Fiche d'identité enregistrée.", "success")
        conn.close()
        return redirect(url_for("client_info"))
    row = conn.execute("SELECT * FROM client_info WHERE client_id=?", (cl["id"],)).fetchone()
    conn.close()
    return render_template("client_info.html", client=cl, info=dict(row) if row else {},
                           clients=database.user_clients(g.user))


@app.route("/admin/clients/<int:cid>/logo", methods=["POST"])
@login_required
@admin_required
def upload_client_logo(cid):
    logo = request.files.get("logo")
    if not logo or not logo.filename:
        flash("Aucun fichier reçu.", "error")
        return redirect(url_for("admin"))
    ext = os.path.splitext(logo.filename)[1].lower()
    if ext not in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"):
        flash("Format non supporté (PNG, JPG, GIF, WebP ou SVG).", "error")
        return redirect(url_for("admin"))
    upload_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads")
    os.makedirs(upload_dir, exist_ok=True)
    filename = "client_%d_logo%s" % (cid, ext)
    logo.save(os.path.join(upload_dir, filename))
    conn = database.get_db()
    conn.execute("UPDATE clients SET logo_path=? WHERE id=?", (os.path.join(upload_dir, filename), cid))
    conn.commit()
    conn.close()
    flash("Logo du client enregistré.", "success")
    return redirect(url_for("admin"))


@app.route("/uploads/<path:filename>")
def uploaded_file(filename):
    upload_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads")
    return send_file(os.path.join(upload_dir, filename))


@app.route("/archive", methods=["POST"])
@login_required
def archive_version():
    cl = current_client()
    if not cl:
        abort(403)
    client, units, actions = _load_export_data()
    buf = io.BytesIO()
    exports.build_pdf(client, units, actions, buf)
    archive_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "archives")
    os.makedirs(archive_dir, exist_ok=True)
    from datetime import datetime
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    filename = "DUERP_%s_%s.pdf" % (client["name"].replace(" ", "_"), stamp)
    with open(os.path.join(archive_dir, filename), "wb") as f:
        f.write(buf.getvalue())
    conn = database.get_db()
    conn.execute(
        "INSERT INTO archives (client_id, filename) VALUES (?,?)", (client["id"], filename))
    conn.commit()
    conn.close()
    flash("Version archivée : %s" % filename, "success")
    return redirect(url_for("home"))


@app.route("/archive/<int:aid>/download")
@login_required
def download_archive(aid):
    cl = current_client()
    if not cl:
        abort(403)
    conn = database.get_db()
    row = conn.execute("SELECT * FROM archives WHERE id=? AND client_id=?", (aid, cl["id"])).fetchone()
    conn.close()
    if not row:
        abort(404)
    archive_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "archives")
    return send_file(os.path.join(archive_dir, row["filename"]), as_attachment=True,
                     download_name=row["filename"])


# ---------------- Assistant IA (Mistral) ----------------

def _client_ai_enabled(cl):
    conn = database.get_db()
    row = conn.execute("SELECT ai_enabled FROM clients WHERE id=?", (cl["id"],)).fetchone()
    conn.close()
    return bool(row and row["ai_enabled"])


@app.route("/client/ai-toggle", methods=["POST"])
@login_required
def client_ai_toggle():
    cl = current_client()
    if not cl:
        abort(403)
    if g.user["role"] != "admin":
        abort(403)
    conn = database.get_db()
    conn.execute("UPDATE clients SET ai_enabled = 1 - ai_enabled WHERE id=?", (cl["id"],))
    conn.commit()
    state = conn.execute("SELECT ai_enabled FROM clients WHERE id=?", (cl["id"],)).fetchone()
    conn.close()
    return jsonify({"ai_enabled": bool(state["ai_enabled"])})


@app.route("/ai/status")
@login_required
def ai_status():
    cl = current_client()
    enabled = _client_ai_enabled(cl) if cl else False
    return jsonify({"configured": ai.is_configured(), "client_enabled": enabled})


@app.route("/ai/suggest-situations", methods=["POST"])
@login_required
def ai_suggest_situations():
    cl = current_client()
    if not cl:
        abort(403)
    if not _client_ai_enabled(cl):
        return jsonify({"error": "L'assistant IA est désactivé pour ce client."}), 503
    data = request.get_json(silent=True) or {}
    unit_name = (data.get("unit_name") or "").strip()
    if not unit_name:
        return jsonify({"error": "Unité de travail manquante."}), 400
    # 1) base de connaissances du client (prioritaire, sans API)
    conn = database.get_db()
    suggestions = []
    seen = set()
    for kw in [w for w in unit_name.replace("-", " ").split() if len(w) > 3]:
        for s in ai.kb_match(conn, cl["id"], kw):
            key = (s["hazard"], s["situation"])
            if key not in seen:
                seen.add(key)
                suggestions.append(s)
        if len(suggestions) >= 6:
            break
    if not suggestions:
        # pas de correspondance par mots-clés : proposer les entrées
        # classiques du client (les plus récemment ajoutées)
        for s in ai.kb_recent(conn, cl["id"], 6):
            key = (s["hazard"], s["situation"])
            if key not in seen:
                seen.add(key)
                suggestions.append(s)
    conn.close()
    # 2) complément Mistral si configurée et activée globalement
    if len(suggestions) < 6 and ai.is_configured():
        try:
            for s in ai.suggest_situations(unit_name)[:6 - len(suggestions)]:
                try:
                    f, g, m = int(s.get("frequency")), int(s.get("severity")), float(s.get("mastery"))
                except (TypeError, ValueError):
                    continue
                if f in FREQ and g in SEV and m in MASTERY:
                    suggestions.append({
                        "hazard": str(s.get("hazard", "")).strip(),
                        "situation": str(s.get("situation", "")).strip(),
                        "frequency": f, "severity": g, "mastery": m,
                        "source": "Mistral AI",
                    })
        except Exception:
            pass
    if not suggestions and not ai.is_configured():
        return jsonify({"error": "Aucune correspondance dans la base de connaissances et clé API Mistral non configurée."}), 503
    return jsonify({"suggestions": suggestions[:6]})


@app.route("/ai/suggest-measures", methods=["POST"])
@login_required
def ai_suggest_measures():
    cl = current_client()
    if not cl:
        abort(403)
    if not _client_ai_enabled(cl):
        return jsonify({"error": "L'assistant IA est désactivé pour ce client."}), 503
    data = request.get_json(silent=True) or {}
    hazard = (data.get("hazard") or "").strip()
    situation = (data.get("situation") or "").strip()
    if not hazard or not situation:
        return jsonify({"error": "Risque et situation requis."}), 400
    result = {"measures": [], "actions": [], "source": "base de connaissances"}
    conn = database.get_db()
    for s in ai.kb_match(conn, cl["id"], hazard + " " + situation):
        result["measures"].extend(s["measures"])
        result["actions"].extend(s["actions"])
    conn.close()
    if not result["measures"] and not result["actions"] and ai.is_configured():
        try:
            api_result = ai.suggest_measures(hazard, situation)
            result = {**api_result, "source": "Mistral AI"}
        except Exception as e:
            return jsonify({"error": "Erreur API Mistral : %s" % e}), 502
    if not result["measures"] and not result["actions"]:
        return jsonify({"error": "Aucune correspondance dans la base de connaissances et clé API Mistral non configurée."}), 503
    return jsonify(result)


@app.route("/admin/clients/<int:cid>/ai-toggle", methods=["POST"])
@login_required
@admin_required
def admin_ai_toggle(cid):
    conn = database.get_db()
    conn.execute("UPDATE clients SET ai_enabled = 1 - ai_enabled WHERE id=?", (cid,))
    conn.commit()
    conn.close()
    return redirect(url_for("admin"))


# ---------------- Base de connaissances ----------------

ALLOWED_KB_EXT = (".xlsx", ".xls", ".csv", ".pdf", ".txt", ".md")


def _kb_import_tabular(path):
    """Importe un fichier tabulaire (xlsx/csv) au format DUERP."""
    entries = []
    import csv as _csv
    rows = None
    if path.lower().endswith((".xlsx", ".xls")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active or wb[wb.sheetnames[0]]
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        with open(path, encoding="utf8", errors="ignore") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                delim = ";" if sample.count(";") >= sample.count(",") else ","
                rows = list(_csv.reader(f, delimiter=delim))
            except Exception:
                return []
    if not rows:
        return []

    def find(header, contains, equals=()):
        for i, h in enumerate(header):
            if h in equals:
                return i
        for i, h in enumerate(header):
            if any(c in h for c in contains):
                return i
        return None

    def detect(header):
        def short_col(h):
            return len(h) <= 30
        cols = {
            "hazard": find(header, ["risque", "danger", "classe de risques", "famille"]),
            "situation": find(header, ["situation"]),
            "measures": find(header, ["mesure"]),
            "actions": find(header, ["action"]),
            "frequency": find(header, ["fréquence", "frequence"], equals=("f",)),
            "severity": find(header, ["gravité", "gravite"], equals=("g",)),
            "mastery": find(header, ["maîtrise", "maitrise"], equals=("m",)),
        }
        # une vraie ligne d'en-tête : situation reconnue ET (risque OU mesure)
        if cols["situation"] is None:
            return None
        if cols["hazard"] is None and cols["measures"] is None:
            return None
        # et la cellule « situation » doit être un intitulé court, pas une phrase de titre
        if len(header[cols["situation"]]) > 40:
            return None
        return cols

    cols = None
    header_idx = 0
    for i, row in enumerate(rows[:10]):
        header = [str(h or "").strip().lower() for h in row]
        cols = detect(header)
        if cols:
            header_idx = i
            break
    if not cols:
        return []
    for row in rows[header_idx + 1:]:
        def g(key):
            i = cols[key]
            return (row[i].strip() if i is not None and i < len(row) else "")
        hazard, situation = g("hazard"), g("situation")
        if not hazard and not situation:
            continue
        try:
            f = int(float(g("frequency") or 2))
            sev = int(float(g("severity") or 2))
            m = float(g("mastery") or 1)
        except ValueError:
            f, sev, m = 2, 2, 1
        f = f if f in FREQ else 2
        sev = sev if sev in SEV else 2
        m = m if m in MASTERY else 1
        measures = "\n".join(x.strip() for x in (g("measures").replace("\r", "").split("\n") + [g("measures")]) if x.strip())
        actions = g("actions").replace("\r", "")
        entries.append((hazard or situation, situation or hazard, f, sev, m, measures, actions))
    return entries


def _kb_import_pdf_text(path):
    """Extrait le texte d'un PDF."""
    try:
        import pymupdf
        doc = pymupdf.open(path)
        text = "\n".join(p.get_text() for p in doc)
        doc.close()
        return text
    except Exception:
        return ""


@app.route("/kb/import", methods=["POST"])
@login_required
def kb_import():
    cl = current_client()
    if not cl:
        abort(403)
    f = request.files.get("file")
    if not f or not f.filename:
        flash("Aucun fichier reçu.", "error")
        return redirect(url_for("kb_list"))
    ext = os.path.splitext(f.filename)[1].lower()
    if ext not in ALLOWED_KB_EXT:
        flash("Format non supporté. Formats acceptés : Excel (.xlsx/.xls), CSV, PDF, texte (.txt/.md).", "error")
        return redirect(url_for("kb_list"))
    tmp = os.path.join(tempfile.gettempdir(), "kb_import_%d%s" % (int(time.time() * 1000), ext))
    f.save(tmp)
    entries = []
    if ext in (".xlsx", ".xls", ".csv"):
        entries = _kb_import_tabular(tmp)
        source = "fichier"
    else:
        text = open(tmp, encoding="utf8", errors="ignore").read() if ext in (".txt", ".md") else _kb_import_pdf_text(tmp)
        if not text.strip():
            flash("Impossible de lire le contenu du fichier.", "error")
            os.remove(tmp)
            return redirect(url_for("kb_list"))
        if ai.is_configured():
            try:
                raw = ai.extract_situations(text)
            except Exception as e:
                flash("Erreur API Mistral : %s" % e, "error")
                os.remove(tmp)
                return redirect(url_for("kb_list"))
            for e in raw:
                try:
                    feq = int(e.get("frequency"))
                    sev = int(e.get("severity"))
                    m = float(e.get("mastery"))
                except (TypeError, ValueError):
                    feq, sev, m = 2, 2, 1
                feq = feq if feq in FREQ else 2
                sev = sev if sev in SEV else 2
                m = m if m in MASTERY else 1
                measures = "\n".join(str(x) for x in (e.get("measures") or []) if str(x).strip())
                acts = []
                for act in (e.get("actions") or []):
                    if isinstance(act, dict) and act.get("text"):
                        acts.append(" | ".join([act["text"], act.get("actor", ""), act.get("due_date", "")]).rstrip(" |"))
                entries.append((str(e.get("hazard", "")).strip(), str(e.get("situation", "")).strip(),
                                feq, sev, m, measures, "\n".join(acts)))
            source = "IA"
        else:
            flash("Pour les PDF et documents texte, la clé API Mistral est nécessaire (MISTRAL_API_KEY). Les fichiers Excel/CSV s'importent sans clé.", "error")
            os.remove(tmp)
            return redirect(url_for("kb_list"))
    os.remove(tmp)
    if not entries:
        flash("Aucune situation dangereuse n'a pu être extraite du fichier.", "error")
        return redirect(url_for("kb_list"))
    conn = database.get_db()
    for hz, st, feq, sev, m, ms, ac in entries[:200]:
        conn.execute(
            "INSERT INTO kb_entries (client_id, hazard, situation, frequency, severity, mastery, measures, actions) "
            "VALUES (?,?,?,?,?,?,?,?)", (cl["id"], hz, st, feq, sev, m, ms, ac))
    conn.commit()
    n = len(entries[:200])
    conn.close()
    flash("%d entrée(s) ajoutée(s) à la base de connaissances (source : %s)." % (n, source), "success")
    return redirect(url_for("kb_list"))


@app.route("/kb")

@app.route("/kb")
@login_required
def kb_list():
    cl = current_client()
    if not cl:
        abort(403)
    conn = database.get_db()
    entries = conn.execute(
        "SELECT * FROM kb_entries WHERE client_id=? ORDER BY id DESC", (cl["id"],)
    ).fetchall()
    conn.close()
    from risk_catalog import RISQUE_CATALOG
    return render_template("kb.html", client=cl, entries=[dict(e) for e in entries],
                           catalog=RISQUE_CATALOG,
                           clients=database.user_clients(g.user))


@app.route("/kb/add", methods=["POST"])
@login_required
def kb_add():
    cl = current_client()
    if not cl:
        abort(403)
    try:
        f = int(request.form.get("frequency", 2))
        g = int(request.form.get("severity", 2))
        m = float(request.form.get("mastery", 1))
    except (TypeError, ValueError):
        abort(400)
    if f not in FREQ or g not in SEV or m not in MASTERY:
        abort(400)
    conn = database.get_db()
    conn.execute(
        "INSERT INTO kb_entries (client_id, hazard, situation, frequency, severity, mastery, measures, actions) "
        "VALUES (?,?,?,?,?,?,?,?)",
        (cl["id"], request.form.get("hazard", "").strip(), request.form.get("situation", "").strip(),
         f, g, m, request.form.get("measures", "").strip(), request.form.get("actions", "").strip()),
    )
    conn.commit()
    conn.close()
    flash("Entrée ajoutée à la base de connaissances.", "success")
    return redirect(url_for("kb_list"))


@app.route("/kb/<int:eid>/delete", methods=["POST"])
@login_required
def kb_delete(eid):
    cl = current_client()
    if not cl:
        abort(403)
    conn = database.get_db()
    conn.execute("DELETE FROM kb_entries WHERE id=? AND client_id=?", (eid, cl["id"]))
    conn.commit()
    conn.close()
    return redirect(url_for("kb_list"))


# ---------------- Exports ----------------

def _client_full(client_id):
    """Client avec logo (chemin absolu) et fiche d'identité."""
    conn = database.get_db()
    row = conn.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    info = conn.execute("SELECT * FROM client_info WHERE client_id=?", (client_id,)).fetchone()
    conn.close()
    c = dict(row) if row else {}
    c["logo_abs"] = c.get("logo_path") or None
    c["info"] = dict(info) if info else {}
    return c


def _last_update(conn, client_id):
    """Date de dernière modification du contenu DUERP d'un client."""
    dates = []
    for table, col, where in (
        ("work_units", "created_at", "client_id"),
        ("risks", "created_at", None),
        ("measures", "created_at", None),
        ("actions", "created_at", None),
        ("action_items", "created_at", "client_id"),
        ("client_info", "updated_at", "client_id"),
    ):
        try:
            if where:
                rows = conn.execute("SELECT %s FROM %s WHERE %s=?" % (col, table, where), (client_id,)).fetchall()
            else:
                rows = conn.execute(
                    "SELECT t.%s FROM %s t JOIN risks r ON r.id=t.risk_id "
                    "JOIN work_units w ON w.id=r.work_unit_id WHERE w.client_id=?" % (col, table),
                    (client_id,)).fetchall()
            dates.extend(r[0] for r in rows if r[0])
        except Exception:
            pass
    return max(dates) if dates else None


def _load_export_data():
    client = current_client()
    if not client:
        abort(403)
    client = _client_full(client["id"])
    units = _load_units(client["id"])
    conn = database.get_db()
    actions = [dict(a) for a in conn.execute(
        "SELECT a.*, w.name AS work_unit FROM action_items a "
        "LEFT JOIN work_units w ON w.id = a.work_unit_id "
        "WHERE a.client_id = ? ORDER BY a.id", (client["id"],)
    ).fetchall()]
    conn.close()
    return client, units, actions


@app.route("/export/pdf")
@login_required
def export_pdf():
    client, units, actions = _load_export_data()
    buf = io.BytesIO()
    exports.build_pdf(client, units, actions, buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="DUERP_%s.pdf" % client["name"].replace(" ", "_"))


@app.route("/export/excel")
@login_required
def export_excel():
    client, units, actions = _load_export_data()
    buf = io.BytesIO()
    exports.build_excel(client, units, actions, buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        download_name="DUERP_%s.xlsx" % client["name"].replace(" ", "_"))


@app.route("/export/action-pdf")
@login_required
def export_action_pdf():
    client, _units, actions = _load_export_data()
    buf = io.BytesIO()
    exports.build_action_pdf(client, actions, buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="Plan_action_%s.pdf" % client["name"].replace(" ", "_"))


@app.route("/export/action-excel")
@login_required
def export_action_excel():
    client, _units, actions = _load_export_data()
    buf = io.BytesIO()
    exports.build_action_excel(client, actions, buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        download_name="Plan_action_%s.xlsx" % client["name"].replace(" ", "_"))


# ---------------- Administration ----------------

@app.route("/admin")
@login_required
@admin_required
def admin():
    conn = database.get_db()
    clients = conn.execute("SELECT * FROM clients ORDER BY name").fetchall()
    users = conn.execute("SELECT * FROM users ORDER BY email").fetchall()
    memberships = conn.execute("SELECT * FROM client_members").fetchall()
    conn.close()
    return render_template("admin.html", clients=clients, users=users, memberships=memberships)


@app.route("/admin/clients/add", methods=["POST"])
@login_required
@admin_required
def admin_add_client():
    name = request.form.get("name", "").strip()
    if name:
        conn = database.get_db()
        conn.execute("INSERT OR IGNORE INTO clients (name) VALUES (?)", (name,))
        conn.commit()
        conn.close()
    return redirect(url_for("admin"))


@app.route("/admin/users/add", methods=["POST"])
@login_required
@admin_required
def admin_add_user():
    email = request.form.get("email", "").strip().lower()
    password = request.form.get("password", "")
    role = request.form.get("role", "user")
    if email and password and role in ("admin", "user"):
        conn = database.get_db()
        try:
            conn.execute(
                "INSERT INTO users (email, password_hash, full_name, role) VALUES (?,?,?,?)",
                (email, generate_password_hash(password), request.form.get("full_name", "").strip(), role),
            )
            conn.commit()
        except Exception:
            flash("Cet email existe déjà.", "error")
        conn.close()
    return redirect(url_for("admin"))


@app.route("/admin/members/add", methods=["POST"])
@login_required
@admin_required
def admin_add_member():
    try:
        user_id = int(request.form.get("user_id"))
        client_id = int(request.form.get("client_id"))
    except (TypeError, ValueError):
        abort(400)
    conn = database.get_db()
    conn.execute("INSERT OR IGNORE INTO client_members (client_id, user_id) VALUES (?,?)", (client_id, user_id))
    conn.commit()
    conn.close()
    return redirect(url_for("admin"))


@app.route("/admin/members/<int:mid>/delete", methods=["POST"])
@login_required
@admin_required
def admin_delete_member(mid):
    conn = database.get_db()
    conn.execute("DELETE FROM client_members WHERE id = ?", (mid,))
    conn.commit()
    conn.close()
    return redirect(url_for("admin"))


if __name__ == "__main__":
    database.init_db()
    database.create_admin()
    app.run(debug=True, host="127.0.0.1", port=int(os.environ.get("PORT", 5000)))
